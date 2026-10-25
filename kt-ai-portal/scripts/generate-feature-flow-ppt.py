#!/usr/bin/env python3
"""Generate feature-flow focused proposal PPT."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

KOREAN_FONT = "Apple SD Gothic Neo"
TITLE_COLOR = RGBColor(0x0D, 0x47, 0x21)
TEXT_COLOR = RGBColor(0x33, 0x33, 0x33)
ACCENT_COLOR = RGBColor(0x1B, 0x5E, 0x20)


def set_font(run, size=Pt(14), bold=False, color=TEXT_COLOR):
    run.font.size = size
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = KOREAN_FONT


ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_DIR = os.path.join(ROOT, "docs", "assets", "final")
OUT_PATH = os.path.join(
    ROOT,
    "proposal-package",
    "국방지능형플랫폼포털_기능흐름_RFP대응_발표자료_32_2단계.pptx",
)

def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    box = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(12.333), Inches(1.5))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = "국방지능형플랫폼포털 기능흐름 RFP 대응 발표자료"
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], size=Pt(40), bold=True, color=TITLE_COLOR)

    sub = slide.shapes.add_textbox(Inches(0.5), Inches(4.0), Inches(12.333), Inches(1.2))
    stf = sub.text_frame
    stf.word_wrap = True
    sp = stf.paragraphs[0]
    sp.text = "사용자 중심 기능 흐름 · 플랫폼 연계 역할 · RFP 21개 SFR 대응 정리\n2026-07-13"
    sp.alignment = PP_ALIGN.CENTER
    set_font(sp.runs[0], size=Pt(20), color=TEXT_COLOR)


def add_bullet_slide(prs, title, bullets, slide_num=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.95)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = TITLE_COLOR
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.6))
    tf = tbox.text_frame
    p = tf.paragraphs[0]
    p.text = title
    set_font(p.runs[0], size=Pt(28), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    footer = slide.shapes.add_textbox(Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.3))
    ftf = footer.text_frame
    ftf.text = "국방지능형플랫폼포털"
    fp = ftf.paragraphs[0]
    fp.alignment = PP_ALIGN.RIGHT
    set_font(fp.runs[0], size=Pt(10), color=RGBColor(0x88, 0x88, 0x88))

    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.25), Inches(12.0), Inches(5.75))
    btf = body.text_frame
    btf.word_wrap = True
    for i, text in enumerate(bullets):
        para = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        para.text = text
        para.level = 0
        set_font(para.runs[0], size=Pt(20), color=TEXT_COLOR)
        para.space_after = Pt(12)


def add_two_column_slide(prs, title, left_title, left_bullets, right_title, right_bullets, slide_num=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.95)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = TITLE_COLOR
    bar.line.fill.background()

    tbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.333), Inches(0.6))
    p = tbox.text_frame.paragraphs[0]
    p.text = title
    set_font(p.runs[0], size=Pt(28), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    footer = slide.shapes.add_textbox(Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.3))
    ftf = footer.text_frame
    ftf.text = "국방지능형플랫폼포털"
    fp = ftf.paragraphs[0]
    fp.alignment = PP_ALIGN.RIGHT
    set_font(fp.runs[0], size=Pt(10), color=RGBColor(0x88, 0x88, 0x88))

    def add_col(col_left, col_top, col_w, col_h, col_title, items):
        box = slide.shapes.add_textbox(col_left, col_top, col_w, col_h)
        tf = box.text_frame
        tf.word_wrap = True
        tp = tf.paragraphs[0]
        tp.text = col_title
        set_font(tp.runs[0], size=Pt(20), bold=True, color=ACCENT_COLOR)
        tp.space_after = Pt(10)
        for i, item in enumerate(items):
            para = tf.add_paragraph()
            para.text = item
            set_font(para.runs[0], size=Pt(16), color=TEXT_COLOR)
            para.space_after = Pt(8)

    add_col(Inches(0.5), Inches(1.2), Inches(5.9), Inches(5.8), left_title, left_bullets)
    add_col(Inches(6.9), Inches(1.2), Inches(5.9), Inches(5.8), right_title, right_bullets)


def add_content_slide(prs, data, slide_num=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Title bar
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.95)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = TITLE_COLOR
    bar.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.6))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"{data['screen_name']} — {data['rfp_ids']}"
    set_font(p.runs[0], size=Pt(24), bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.4), Inches(7.1), Inches(12.5), Inches(0.3))
    ftf = footer.text_frame
    ftf.text = "국방지능형플랫폼포털"
    fp = ftf.paragraphs[0]
    fp.alignment = PP_ALIGN.RIGHT
    set_font(fp.runs[0], size=Pt(10), color=RGBColor(0x88, 0x88, 0x88))

    left = Inches(0.4)
    top = Inches(1.05)
    width = Inches(5.4)
    height = Inches(6.0)
    content_box = slide.shapes.add_textbox(left, top, width, height)
    ctf = content_box.text_frame
    ctf.word_wrap = True

    sections = [
        ("RFP ID", data["rfp_ids"]),
        ("요구사항명", data["req_name"]),
        ("요구내용", data["req_content"]),
        ("대응 화면", f"{data['screen_name']} ({data['screen_route']})"),
        ("구현 기능", "\n".join(f"• {f}" for f in data["features"])),
        ("구현 수준", data["impl_level"]),
        ("후속 연동 범위", data["follow_up"]),
    ]

    for i, (label, value) in enumerate(sections):
        para = ctf.paragraphs[0] if i == 0 else ctf.add_paragraph()
        para.text = f"{label}: "
        set_font(para.runs[0], size=Pt(14), bold=True, color=ACCENT_COLOR)
        para.space_after = Pt(2)

        val_para = ctf.add_paragraph()
        val_para.text = value
        set_font(val_para.runs[0], size=Pt(13), color=TEXT_COLOR)
        val_para.space_after = Pt(8)

    img_path = os.path.join(IMG_DIR, data["image_file"])
    if os.path.exists(img_path):
        img_left = Inches(6.0)
        img_top = Inches(1.05)
        img_max_height = Inches(6.0)
        img_max_width = Inches(7.0)
        pic = slide.shapes.add_picture(img_path, img_left, img_top, height=img_max_height)
        if pic.width > img_max_width:
            ratio = img_max_width / pic.width
            pic.width = img_max_width
            pic.height = int(pic.height * ratio)
    else:
        missing = slide.shapes.add_textbox(Inches(6.0), Inches(3.0), Inches(7.0), Inches(1.0))
        mtf = missing.text_frame
        mtf.text = f"캡처 이미지 없음:\n{data['image_file']}"
        set_font(mtf.paragraphs[0].runs[0], size=Pt(14), color=RGBColor(0x99, 0x00, 0x00))


def add_section_divider(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(12.333), Inches(1.2))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    set_font(p.runs[0], size=Pt(36), bold=True, color=TITLE_COLOR)
    if subtitle:
        sp = tf.add_paragraph()
        sp.text = subtitle
        sp.alignment = PP_ALIGN.CENTER
        set_font(sp.runs[0], size=Pt(20), color=TEXT_COLOR)
        sp.space_before = Pt(12)


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    add_title_slide(prs)

    add_bullet_slide(
        prs,
        "포털은 ‘통합 진입점’입니다",
        [
            "포털은 데이터 플랫폼, AI 플랫폼, 지능형 클라우드/CMP를 직접 대체하지 않습니다.",
            "사용자가 하나의 접점에서 각 플랫폼을 안전하게 접근·신청·조회·승인·연계할 수 있도록 지원합니다.",
            "데이터 플랫폼 = 데이터를 수집·정리·가공하는 곳",
            "AI 플랫폼 = 모델을 개발·학습·배포하는 곳",
            "지능형 클라우드/CMP = GPU·서버·스토리지 자원을 제공하는 곳",
            "포털은 이 세 영역을 하나의 화면에서 신청하고 조회하고 승인할 수 있게 연결합니다.",
        ],
    )

    add_bullet_slide(
        prs,
        "전체 체계 개요",
        [
            "국방지능형플랫폼포털 — 단일 접점 · 권한 통제 · 업무 오케스트레이션",
            "데이터 플랫폼 — 데이터 수집/저장/가공/분석, 카탈로그, OpenSearch",
            "AI 플랫폼 — 모델 개발/학습/배포, 작업환경, 모델 레지스트리",
            "지능형 클라우드/CMP — 인프라/컴퓨팅/스토리지 자원, 과금/쿼터",
            "인증/권한 — SSO, 사용자/조직/역할/테넌트 관리",
            "보안/K-RMF — 반입·반출 승인, 증빙 관리, 감사 로그",
        ],
    )

    add_two_column_slide(
        prs,
        "사용자 유형별 사용 흐름",
        "일반 사용자",
        [
            "접속 → 로그인 → 대시보드 확인",
            "통합검색/카탈로그 탐색",
            "필요 시 프로젝트/반출/크레딧 업무 신청",
        ],
        "관리자/담당자",
        [
            "보안 관리자: 반입·반출 승인, K-RMF 증빙, 감사",
            "군별 관리자: 테넌트 쿼터/사용량/권한 통제",
            "사업관리자: 요구사항 대응현황, API 연계현황, 시나리오 추적",
        ],
    )

    add_bullet_slide(
        prs,
        "전체 메뉴 구조",
        [
            "포털 공통 — 통합 대시보드, 통합검색, 서비스/도구 카탈로그, 제안 게시판",
            "프로젝트 운영 — 프로젝트/워크스페이스, 반입·반출 승인, 크레딧 관리",
            "관리/보안 — K-RMF 증빙관리, 관리자 콘솔",
            "사용자/조직 — 마이페이지, 환경설정, 육해공군 관리",
            "관리자 전용 — 요구사항 대응현황, 플랫폼 API 연계현황, 업무흐름/연계 시나리오",
        ],
    )

    add_bullet_slide(
        prs,
        "RFP 21개 SFR 대응 총괄",
        [
            "SFR-001 메인 포털 — 통합 대시보드",
            "SFR-002 로그인/로그아웃, SFR-003~006 관리자/권한/정책/메뉴",
            "SFR-007~008 작업환경 생성, SFR-009~011 반입·반출 승인",
            "SFR-012 플랫폼 연동, SFR-013 제안게시판",
            "SFR-014 카탈로그, SFR-015 통합검색",
            "SFR-016~018 프로젝트/크레딧/워크스페이스, SFR-019~021 마이페이지/환경설정/육해공군",
        ],
    )

    add_section_divider(prs, "SFR별 기능 흐름 및 화면", "RFP 요구사항과 구현 화면을 매핑합니다")

    feature_slides = [
        {
            "screen_name": "통합 대시보드",
            "screen_route": "/",
            "rfp_ids": "SFR-001, SFR-012",
            "req_name": "메인 포털 / 플랫폼 연동",
            "req_content": "플랫폼 전체 현황을 요약하고 공지사항, Quick Links, 시스템 상태를 제공합니다.",
            "features": ["요약 카드", "공지/Quick Links", "통합검색 진입", "시스템 상태 패널"],
            "impl_level": "화면 구현 완료",
            "follow_up": "실시간 플랫폼 상태 API, 개인화 위젯",
            "image_file": "01-portal-dashboard.png",
        },
        {
            "screen_name": "로그인 / 인증",
            "screen_route": "/",
            "rfp_ids": "SFR-002",
            "req_name": "메인 포털 로그인/로그아웃",
            "req_content": "접근코드 또는 Google 계정으로 로그인하고, Keycloak 기반 인증을 준비합니다.",
            "features": ["접근코드 로그인", "Google 로그인", "Firebase Auth 연동", "Keycloak OIDC 연동 준비"],
            "impl_level": "화면 구현 완료 (Firebase Auth 연동, Keycloak OIDC 연동 확인)",
            "follow_up": "운영 SSO, MFA, PKI/MPKI 연동",
            "image_file": "09-backstage-signin.png",
        },
        {
            "screen_name": "관리자 콘솔",
            "screen_route": "/admin-console",
            "rfp_ids": "SFR-003~006",
            "req_name": "사용자/권한/정책/메뉴 관리",
            "req_content": "관리자가 사용자, 역할, 메뉴, 정책, 감사로그를 중앙에서 관리합니다.",
            "features": ["사용자 등록/승인", "역할/그룹/권한", "메뉴/프로그램 관리", "로그인/비밀번호/세션 정책", "감사로그"],
            "impl_level": "화면 구현 완료 (관리자 전용, 샘플 데이터 기반)",
            "follow_up": "Keycloak Admin API, 조직/부서 DB, 정책 엔진 연동",
            "image_file": "07-admin-console-user-management.png",
        },
        {
            "screen_name": "서비스/도구 카탈로그",
            "screen_route": "/service-catalog",
            "rfp_ids": "SFR-007, SFR-008, SFR-014",
            "req_name": "AI/데이터 작업환경 생성 요청 / 카탈로그 조회",
            "req_content": "AI 개발에 필요한 분석/개발/데이터/시각화/라벨링 도구를 조회·선택·요청합니다.",
            "features": ["도구 카테고리 탭", "도구 목록/상세", "도구 선택/직접 입력 요청"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "OpenMetadata/Docker/PyPI 실제 연동, 도구 요청 workflow",
            "image_file": "10-service-catalog.png",
        },
        {
            "screen_name": "반입·반출 승인",
            "screen_route": "/export-approval",
            "rfp_ids": "SFR-009~011, SFR-012",
            "req_name": "반입·반출 관리 / 플랫폼 연동",
            "req_content": "데이터 반입/반출에 대한 신청·승인·보안 검토 workflow를 지원합니다.",
            "features": ["반입/반출 신청 목록", "신청 등록 및 보안등급 설정", "승인/반려 workflow", "이력 관리"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "실제 승인 workflow, 파일 저장소, 보안 검토 API, 감사로그 DB",
            "image_file": "04-export-approval.png",
        },
        {
            "screen_name": "통합검색",
            "screen_route": "/integrated-search",
            "rfp_ids": "SFR-014, SFR-015",
            "req_name": "카탈로그 조회 / 통합검색",
            "req_content": "데이터셋·모델·도구·문서를 보안등급과 유형 중심으로 통합 검색합니다.",
            "features": ["통합 검색어 입력", "유형/보안등급/소스/태그 필터", "메타데이터 상세 패널"],
            "impl_level": "화면 구현 완료 (OpenSearch proxy 연동, 샘플 데이터 fallback)",
            "follow_up": "운영 OpenSearch 인스턴스, 권한 기반 결과 필터",
            "image_file": "02-integrated-search.png",
        },
        {
            "screen_name": "프로젝트/워크스페이스",
            "screen_route": "/project-workspace",
            "rfp_ids": "SFR-016, SFR-018",
            "req_name": "프로젝트 관리 / 워크스페이스 관리 요청",
            "req_content": "AI 연구/개발 프로젝트를 등록·조회하고 워크스페이스·자원을 신청·관리합니다.",
            "features": ["프로젝트 목록/상세", "멤버 및 자원 관리", "워크스페이스 생성 신청", "승인 이력"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "프로젝트 DB/API, Keycloak 그룹, CMP/Kubernetes 워크스페이스 연동",
            "image_file": "03-project-workspace.png",
        },
        {
            "screen_name": "크레딧 관리",
            "screen_route": "/credit-manager",
            "rfp_ids": "SFR-016, SFR-017, SFR-018",
            "req_name": "크레딧 신청 및 조회 / 워크스페이스 관리",
            "req_content": "프로젝트별 크레딧을 할당하고 GPU/CPU/스토리지 사용량을 모니터링합니다.",
            "features": ["프로젝트별 크레딧 잔액/사용량", "충전/사용 이력", "자원별 단가", "부족 알림"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "AI/큐브/클우드 사용량 수집 API, 크레딧 차감 enforcement",
            "image_file": "05-credit-manager.png",
        },
        {
            "screen_name": "K-RMF 증빙관리",
            "screen_route": "/krmf-evidence",
            "rfp_ids": "SFR-002, SFR-004",
            "req_name": "로그인/로그아웃 / 정책 및 이력 관리",
            "req_content": "K-RMF 평가 항목별 증빙 자료를 등록·조회·관리하고 감사 이력과 연계합니다.",
            "features": ["통제항목 목록", "증빙 자료 등록/조회", "보안 미흡사항 조치관리", "통합진척률 대시보드"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "K-RMF 평가 DB, 증빙 파일 저장소, 감사이력 연동",
            "image_file": "06-krmf-evidence.png",
        },
        {
            "screen_name": "마이페이지 / 환경설정 / 육해공군 관리",
            "screen_route": "/my-page, /environment-settings, /tenant-management",
            "rfp_ids": "SFR-019~021",
            "req_name": "마이페이지 / 환경설정 / 육해공군 관리자 권한 관리",
            "req_content": "사용자 개인 관리, 운영 설정, 군별 테넌트·쿼터·사용량 관리를 지원합니다.",
            "features": ["프로필/세션/API Key", "시스템/보안/자원 설정", "군별 테넌트 개요/위임/쿼터/사용량"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "Keycloak/DB, 설정 DB, Kubernetes/OpenStack/CMP quota 연동",
            "image_file": "14-tenant-management.png",
        },
        {
            "screen_name": "제안 게시판",
            "screen_route": "/proposal-board",
            "rfp_ids": "SFR-013",
            "req_name": "제안게시판",
            "req_content": "사용자가 AI·데이터·도구·정책 아이디어를 제안하고 관리자가 검토합니다.",
            "features": ["제안 목록/상세", "제안 등록", "공감/추천/관심/댓글", "관리자 검토"],
            "impl_level": "화면 구현 완료 (샘플 데이터 기반)",
            "follow_up": "제안 DB, 알림/메일, 승인 workflow",
            "image_file": "11-proposal-board.png",
        },
    ]
    for data in feature_slides:
        add_content_slide(prs, data)

    add_section_divider(prs, "관리자 추적·연계 관리 기능")

    admin_slides = [
        {
            "screen_name": "관리자 콘솔 – 플랫폼 운영도구",
            "screen_route": "/admin-console",
            "rfp_ids": "SFR-003~006",
            "req_name": "플랫폼 운영도구 관리자 전용 링크",
            "req_content": "OpenMetadata, OpenSearch 대시보드, Keycloak 관리 등 외부 운영도구를 관리자 전용으로 집중 제공합니다.",
            "features": ["OpenMetadata 메타데이터 관리", "OpenSearch 대시보드 모니터링", "Keycloak 인증·인가 관리", "관리자 전용 접근 제어"],
            "impl_level": "화면 구현 완료 (관리자 전용, 외부 도구 링크)",
            "follow_up": "환경별 외부 도구 URL 설정, SSO 연동",
            "image_file": "09-admin-console-platform-ops-tools.png",
        },
        {
            "screen_name": "요구사항 대응현황",
            "screen_route": "/requirement-coverage",
            "rfp_ids": "SFR-001~SFR-021 전체 추적",
            "req_name": "RFP 163개 기능목록 대응 현황",
            "req_content": "메뉴/SFR/기능 단위로 구현 현황을 체계적으로 조회하고 CSV로 추출합니다.",
            "features": ["메뉴 대응현황", "SFR 대응현황", "기능 대응현황", "변경이력/질의사항", "CSV 다운로드"],
            "impl_level": "화면 구현 완료 (관리자 전용, 샘플 데이터 기반)",
            "follow_up": "RFP 확정 후 데이터 갱신, 사업관리자용 추적 도구로 활용",
            "image_file": "17-requirement-coverage.png",
        },
        {
            "screen_name": "플랫폼 API 연계현황",
            "screen_route": "/platform-api-integration",
            "rfp_ids": "SFR-012 플랫폼 연동",
            "req_name": "포털-플랫폼 연계 후보 관리",
            "req_content": "데이터/AI/클리우드/인증/보안 플랫폼 간 연계 후보를 API/이벤트/파일 단위로 정의합니다.",
            "features": ["연계 후보 목록", "API/이벤트/파일 연계 구분", "연계 플랫폼/메뉴/SFR 매핑", "CSV 다운로드"],
            "impl_level": "화면 구현 완료 (관리자 전용, 설계안/후보 목록 기반)",
            "follow_up": "각 플랫폼 담당 팀과 인터페이스 정의서 확정 및 실제 연동 개발",
            "image_file": "18-platform-api-integration.png",
        },
        {
            "screen_name": "업무흐름/연계 시나리오",
            "screen_route": "/workflow-scenarios",
            "rfp_ids": "SFR-001, SFR-012, SFR-014, SFR-016, SFR-017, SFR-018, SFR-021",
            "req_name": "사용자 여정 및 플랫폼 연계 흐름",
            "req_content": "데이터 엔지니어부터 군별 관리자까지 각 사용자가 포털과 플랫폼을 어떻게 사용하는지 시각화합니다.",
            "features": ["포털 역할 요약", "7개 사용자 여정", "7개 업무흐름 시나리오", "단계별 플랫폼/API/RFP 매핑"],
            "impl_level": "화면 구현 완료 (관리자 전용, 시나리오 데이터 기반)",
            "follow_up": "실제 업무흐름 반영, 특정 시나리오 직접 링크",
            "image_file": "19-workflow-scenarios.png",
        },
    ]
    for data in admin_slides:
        add_content_slide(prs, data)

    add_section_divider(prs, "대표 업무흐름")

    flows = [
        (
            "대표 흐름 1 — 데이터셋 탐색 및 활용 신청",
            [
                "1. 통합검색에서 데이터셋/모델/도구 탐색",
                "2. 보안등급과 메타데이터 확인",
                "3. 프로젝트/워크스페이스에서 활용 신청",
                "4. 데이터 플랫폼 작업환경에서 분석·가공",
                "5. 결과물은 반입·반출 승인을 거쳐 반출",
            ],
        ),
        (
            "대표 흐름 2 — AI 작업환경 생성 및 모델 개발",
            [
                "1. 서비스/도구 카탈로그에서 AI 도구 선택",
                "2. 프로젝트/워크스페이스에 자원/도구 신청",
                "3. 지능형 클라우드/CMP에서 자원 할당",
                "4. AI 플랫폼에서 학습·평가·배포",
                "5. 크레딧 관리에서 사용량과 잔여량 모니터링",
            ],
        ),
        (
            "대표 흐름 3 — 반입·반출 승인 및 감사 추적",
            [
                "1. 반입·반출 승인 메뉴에서 신청 등록",
                "2. 보안등급과 승인권자 자동 분기",
                "3. 보안 관리자가 승인/반려 처리",
                "4. 승인 이력이 감사로그로 축적",
                "5. K-RMF 증빙관리에서 평가 자료로 활용",
            ],
        ),
        (
            "대표 흐름 4 — 크레딧 기반 자원 사용량 관리",
            [
                "1. 프로젝트별 크레딧 잔액 확인",
                "2. GPU/CPU/스토리지 사용량 모니터링",
                "3. 잔여량 부족 시 추가 충전 신청",
                "4. 군별 관리자가 쿼터와 사용량 조회",
                "5. 사용량 기반 자원 회수/재배치 결정",
            ],
        ),
    ]
    for title, bullets in flows:
        add_bullet_slide(prs, title, bullets)

    add_section_divider(prs, "포털과 플랫폼의 연계 역할")

    add_bullet_slide(
        prs,
        "포털 — 데이터 플랫폼 연계",
        [
            "데이터 플랫폼은 데이터 수집·저장·가공·분석을 담당합니다.",
            "포털은 데이터셋/도구 검색, 반입·반출 신청, 프로젝트 연계를 오케스트레이션합니다.",
            "데이터 작업환경 조회, 메타데이터 검색, 데이터 반출 파일 전송을 연계합니다.",
            "실제 데이터 처리는 데이터 플랫폼에서 수행됩니다.",
        ],
    )

    add_bullet_slide(
        prs,
        "포털 — AI 플랫폼 연계",
        [
            "AI 플랫폼은 모델 개발·학습·배포를 담당합니다.",
            "포털은 AI 작업환경/도구 신청, 프로젝트 연계, 성과물 반출 승인을 관리합니다.",
            "AI 모델 조회, 작업환경 상태, 학습/평가 이력을 포털에서 조회합니다.",
            "실제 모델 학습과 배포는 AI 플랫폼에서 수행됩니다.",
        ],
    )

    add_bullet_slide(
        prs,
        "포털 — 지능형 클라우드/CMP 연계",
        [
            "지능형 클라우드/CMP는 GPU·서버·스토리지 자원과 쿼터를 담당합니다.",
            "포털은 프로젝트/워크스페이스 생성 신청, 크레딧 할당, 사용량 모니터링을 관리합니다.",
            "군별 쿼터/사용량/권한 통제를 포털에서 조회하고 위임합니다.",
            "실제 자원 생성과 과금은 클라우드/CMP에서 수행됩니다.",
        ],
    )

    add_bullet_slide(
        prs,
        "현재 구현 수준과 후속 연동 범위",
        [
            "화면 구현: 19개 핵심 화면 UI/UX 구현 완료",
            "관리자 전용 기능: 요구사항 대응현황, 플랫폼 API 연계현황, 업무흐름/연계 시나리오 포함",
            "연동 현황: OpenSearch proxy, OpenMetadata UI, Keycloak OIDC 연동 확인",
            "데이터 기반: 사용자/프로젝트/반출/크레딧/카탈로그 등 샘플 데이터 기반 화면 검증 완료",
            "후속 단계: 실제 DB/API/Keycloak Admin/Kubernetes/OpenStack/CMP/파일 저장소 연동 및 감사로그 DB 구축",
        ],
    )

    add_bullet_slide(
        prs,
        "기대효과 및 마무리",
        [
            "RFP 21개 SFR과 163개 기능목록을 화면과 명확히 매핑",
            "사용자 유형별 업무흐름을 시각적으로 제시",
            "포털과 데이터/AI/클라우드 플랫폼의 역할 경계를 명확히 구분",
            "구현 수준과 후속 연동 범위를 투명하게 공유",
            "제안 발표 및 사업 착수 후 상세설계의 기반 자료로 활용",
        ],
    )

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    prs.save(OUT_PATH)
    print(f"Saved feature-flow PPT: {OUT_PATH}")


if __name__ == "__main__":
    main()
