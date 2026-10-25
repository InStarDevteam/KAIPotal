#!/usr/bin/env python3
"""Build the review workbook from the checked-in report and unchanged FP source."""
import csv
import json
import re
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/requirements-review-20261003'
REPORT = OUT / 'REQUIREMENTS_GAP_REVIEW.md'
SOURCE = ROOT.parent / 'prompt/final-completion-20260929/FP_요구사항_원문기준.csv'
sections = {}
heading = ''
for line in REPORT.read_text().splitlines():
    if line.startswith('## '):
        heading = line[3:]
    elif line.startswith('|'):
        cells = [cell.strip() for cell in line.strip('|').split('|')]
        if all(re.fullmatch(r'[: -]+', cell) for cell in cells):
            continue
        sections.setdefault(heading, []).append(cells)

source_rows = list(csv.DictReader(SOURCE.open(encoding='utf-8-sig')))
assert len(source_rows) == 531 and len({row['FP_ID'] for row in source_rows}) == 531
wb = Workbook()
ws = wb.active
ws.title = '검토안내'
notes = [
    ['항목', '내용'],
    ['기준일', '2026-10-03'],
    ['검토 범위', 'SFR 21영역·FP 원문 531행의 추적, 포털 코드/실제 API/배포 및 UI 점검'],
    ['원문 경계', '견적서 FP 목록과 내부 RFP 경계표 기준. 최종 RFP 전문/승인 기준선은 사업관리팀 확인 필요'],
    ['판정 경계', '원문 531행을 보존·연결한다. 531개 각각의 최신 충족 판정을 완료하거나 현재 완료율을 재산정한 파일이 아니다'],
    ['협의 상태', '외부 팀 합의 없음. 담당 팀은 역할 제안. 실제 담당자·승인자·기한은 미지정'],
    ['구현/연동', '포털 원장 저장, 플랫폼 실행, 기관 인증, 검수 합격은 다른 완료 기준'],
    ['검토 문서', 'REQUIREMENTS_GAP_REVIEW.md'],
    ['실제 배포', 'RELEASE_REPORT.md 및 evidence의 최종 결과 참조'],
    ['비밀정보', '토큰/비밀번호/실개인정보를 이 파일에 넣지 않는다'],
]
for row in notes:
    ws.append(row)

def sheet(name, rows, widths=None, extras=None):
    target = wb.create_sheet(name)
    for index, row in enumerate(rows):
        target.append(row + (extras if index == 0 else [''] * len(extras)) if extras else row)
    if widths:
        for index, width in enumerate(widths, 1):
            target.column_dimensions[get_column_letter(index)].width = width
    return target

req = sections['요구영역별 미비사항']
gaps = sections['미비사항과 처리 구분']
ux = sections['UI/UX 일관성과 편의성']
teams = sections['팀별 협의 의제와 요청 산출물']
assert len(req) == 22 and len(gaps) == 16 and len(ux) == 15 and len(teams) == 13
sheet('요구영역21', req, [24, 44, 70, 40, 56])
sheet('미비사항15', gaps, [12, 10, 52, 92], ['담당자', '목표기한', '결정근거', '처리상태'])
sheet('UIUX14', ux, [12, 70, 100])
sheet('팀협의12', teams, [34, 35, 85, 80], ['담당자', '승인자', '확정기한', '합의여부', '다음조치'])
source_headers = list(source_rows[0])
sheet('FP원문531', [source_headers] + [[row[key] for key in source_headers] for row in source_rows], [18, 12, 12, 26, 30, 64, 95, 12, 16, 16, 22, 28, 34, 65, 28, 35])
areas = {
    '메인 포털': ('SFR-001', 'G03 G07 G08'),
    '로그인/인증': ('SFR-002', 'G01 G02'),
    '사용자 관리': ('SFR-003', 'G01 G02'),
    '권한/RBAC 관리': ('SFR-005', 'G02 G03'),
    '정책 관리': ('SFR-004 SFR-020', 'G01 G11'),
    '메뉴/프로그램 관리': ('SFR-006', 'G03 G08'),
    '통합 카탈로그/검색': ('SFR-014 SFR-015', 'G07 G10'),
    '서비스/도구 카탈로그': ('SFR-007 SFR-008 SFR-014', 'G04 G07'),
    '프로젝트 관리': ('SFR-016', 'G03 G04 G09'),
    '워크스페이스/ 리소스 관리': ('SFR-018', 'G04 G12'),
    '반입/반출 관리': ('SFR-009 SFR-010 SFR-011', 'G05 G06 G10'),
    '크레딧/사용량 관리': ('SFR-017', 'G06 G12'),
    '마이페이지': ('SFR-019', 'G01 G09'),
    '환경설정': ('SFR-020', 'G01 G08'),
    '이력/감사로그': ('SFR-004', 'G11 G15'),
    '제안 게시판': ('SFR-013', 'G08 G09'),
    '육해공군 관리자 권한 관리': ('SFR-021', 'G02 G03 G12'),
    '플랫폼 연동': ('SFR-012', 'G03 G04 G10'),
    'K-RMF 관리지원': ('최신 고객 지시·K-RMF 지원', 'G05 G13'),
    '데이터 관리': ('관련 업무의 저장/외부 참조', 'G03 G07 G13 G14'),
}
trace = [['FP_ID', 'PDF페이지', '원문 영역', '원문 단위프로세스', '관련 요구영역(검토 연결)', '관련 미비 ID(영역 후보)', '판정 범위']]
for row in source_rows:
    area = row['어플리케이션_정리']
    sfr, ids = areas[area]
    trace.append([row['FP_ID'], row['PDF페이지'], area, row['단위프로세스_정리'], sfr, ids, '영역 기반 추적 연결; 개별 FP 미충족 판정 아님'])
sheet('FP검토연결531', trace, [18, 12, 30, 75, 44, 28, 60])

for sheet_index, target in enumerate(wb.worksheets):
    target.freeze_panes = 'A2'
    target.sheet_view.zoomScale = 85
    target.auto_filter.ref = target.dimensions
    target.row_dimensions[1].height = 30
    for cell in target[1]:
        cell.font = Font(name='맑은 고딕', bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='243651')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    for row in target.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name='맑은 고딕', size=10)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        target.row_dimensions[row[0].row].height = 68 if target.title not in ['FP원문531', 'FP검토연결531'] else 92
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill('solid', fgColor='F1F5FA')
    if target.title == '검토안내':
        target.column_dimensions['A'].width = 26
        target.column_dimensions['B'].width = 110
    table = Table(displayName=f'ReviewTable{sheet_index}', ref=target.dimensions)
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    target.add_table(table)

filename = OUT / '요구사항_UIUX_미비사항_협의목록_20261003.xlsx'
wb.save(filename)
check = load_workbook(filename, read_only=True)
assert check['FP원문531'].max_row == 532
assert check['FP검토연결531'].max_row == 532
assert len(check.sheetnames) == 7
for name, rows in [('요구영역21', req), ('미비사항15', gaps), ('UIUX14', ux), ('팀협의12', teams)]:
    with (OUT / f'{name}.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        csv.writer(stream).writerows(rows)
manifest = {'sourceRows': 531, 'uniqueFpIds': 531, 'requirements': 21, 'gaps': 15, 'uxItems': 14, 'teamAgendas': 12, 'sheets': check.sheetnames, 'individualFpFulfillmentReassessed': False, 'externalTeamAgreements': 0}
(OUT / 'workbook-verification.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(manifest, ensure_ascii=False))
